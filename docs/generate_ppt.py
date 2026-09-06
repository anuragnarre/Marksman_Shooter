import collections
import collections.abc
import sys
import os

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
except ImportError:
    print("python-pptx not installed. Installing now...")
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-pptx"])
    from pptx import Presentation
    from pptx.util import Inches, Pt

prs = Presentation()

# Slide 1: Title Slide
slide = prs.slides.add_slide(prs.slide_layouts[0])
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = "MARKSMAN - Master Improvement Plan"
subtitle.text = "All-in-One Platform: Range Operations · Player Management · Coaching\nSeptember 2026"

def add_bullet_slide(title_text, bullets):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = title_text
    body_shape = slide.placeholders[1]
    tf = body_shape.text_frame
    tf.text = bullets[0]
    for bullet in bullets[1:]:
        p = tf.add_paragraph()
        p.text = bullet
        p.level = 0
    return slide

add_bullet_slide("Current State Audit", [
    "✅ Existing: Session CRUD, AI Coach, Ballistics/Equipment stubs, WebSockets.",
    "⚠️ Gaps: No lane/booking management, no member billing, no competition management.",
    "🎯 Goal: Transform into a complete SaaS for ranges, coaches, and shooters."
])

add_bullet_slide("Category 1: API Hardening & Tech Debt", [
    "Crucial foundation before scaling.",
    "• DTO Validation: Enforce strict typing on Equipment, Ranges, Ballistics.",
    "• Security & Limits: Rate limiting, Helmet middleware, strict CORS.",
    "• Performance: Database indexing, pagination, background job queues (Redis).",
    "• Data Integrity: Soft delete consistency and API documentation (Swagger)."
])

add_bullet_slide("Category 2: Range Management Module", [
    "The biggest missing pillar for Range Operators.",
    "• Lane & Bay Management: Real-time lane status, assignments, drag-and-drop UI.",
    "• Bookings: Time slot generation, waitlists, recurring bookings, QR check-ins.",
    "• Membership System: Tiers, digital cards, guest passes, automated billing.",
    "• Range Operations: Dashboard with utilization heatmaps, environmental data sync."
])

add_bullet_slide("Category 3: Player / Shooter Module", [
    "Enhancing the individual athlete experience.",
    "• Comprehensive Profiles: Medical notes, stances, biometrics.",
    "• Achievements & PBs: Automated tracking of personal bests, milestone badges.",
    "• Internal Ranking: ELO-style ratings per discipline and weapon type.",
    "• Training Streaks & Goals: Progress tracking vs. target scores."
])

add_bullet_slide("Category 4: Coach Module", [
    "Empowering coaches to manage squads effectively.",
    "• Squad Management: Group sessions, bulk AI analysis, comparative tables.",
    "• Drill & Template Library: Standardized plans, assignable training events.",
    "• Live Monitoring: Multi-shooter live views, 'Pause & Correct' signals.",
    "• Reporting: Automated PDF progress reports and competition readiness."
])

add_bullet_slide("Category 5: Competition Management", [
    "Turning the platform into a tournament hub.",
    "• Setup & Registration: Online entries, categories, fee collection.",
    "• Live Scoring: Public leaderboards, Kiosk mode, Finals elimination mode.",
    "• Analytics: Performance under pressure (Competition vs. Training delta)."
])

add_bullet_slide("Category 6 & 7: Armory & Safety", [
    "• Equipment: Full weapon registry, service logs, accuracy baselines.",
    "• Ammo Tracking: Inventory deduction, reloading logs, lot performance.",
    "• Safety Rules Engine: Digital waivers, mandatory quizzes.",
    "• Incident Reporting: RSO alerts, photo uploads, compliance audit trails."
])

add_bullet_slide("Implementation Strategy", [
    "1. API Hardening (Blocks all other modules).",
    "2. Range Management (Unlocks B2B revenue).",
    "3. Shooter & Coach Enhancements (Improves retention).",
    "4. Competition Module (Growth & engagement).",
    "5. Hardware Integrations (Premium features)."
])

out_path = os.path.join(os.getcwd(), "Marksman_Master_Plan_Presentation.pptx")
prs.save(out_path)
print(f"Presentation saved to {out_path}")
